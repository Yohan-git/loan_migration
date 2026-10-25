"""Excel report for one scenario: the matrices used and the portfolio vector quarter by quarter.

Sheets, per rating model:
    <model> matrices    annual matrix of years 1, 2, 3+ and the quarterly matrix applied in the projection
    <model> portfolio   amount per class at the end of each quarter, cumulative repayments, total

Figures are computed in Python (`migration.project_portfolio`); totals are simple SUM formulas.

    export_workbook("migration_report.xlsx", matrices, portfolio, "adverse", reference_date=REFERENCE_DATE)
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

from matrices import DEFAULT_SCENARIO, QUARTERS_PER_YEAR
from migration import project_portfolio

FONT = "Arial"
HEADER_FILL = PatternFill("solid", start_color="D9E1F2", end_color="D9E1F2")
DIAGONAL_FILL = PatternFill("solid", start_color="DDEBF7", end_color="DDEBF7")
DEFAULT_FILL = PatternFill("solid", start_color="FCE4D6", end_color="FCE4D6")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
AMOUNT = '#,##0;-#,##0;"-"'
REPAYMENT = {"bullet": "loans are repaid in full at the end of their maturity quarter",
             "linear": "loans are repaid on a straight line down to 0 at their maturity quarter"}


def _font(bold=False, size=10, italic=False, color="000000"):
    return Font(name=FONT, bold=bold, size=size, italic=italic, color=color)


def _title(ws, title, note):
    ws["A1"] = title
    ws["A1"].font = _font(bold=True, size=13)
    ws["A2"] = note
    ws["A2"].font = _font(italic=True, color="595959")
    ws.sheet_view.showGridLines = False


def _header(cell, value):
    cell.value = value
    cell.font = _font(bold=True)
    cell.fill = HEADER_FILL
    cell.border = BORDER
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _value(cell, value, number_format, bold=False, fill=None):
    cell.value = value
    cell.number_format = number_format
    cell.font = _font(bold=bold)
    cell.border = BORDER
    if fill:
        cell.fill = fill


def _quarters(year, last_year):
    q0 = (year - 1) * QUARTERS_PER_YEAR + 1
    return f"quarters {q0}+" if year >= last_year else f"quarters {q0}-{q0 + QUARTERS_PER_YEAR - 1}"


def _write_matrix(ws, top, left, title, df, number_format):
    """Matrix with a title, class headers and a Total column (SUM). Diagonal and default column shaded."""
    n = len(df)
    ws.cell(top, left, title).font = _font(bold=True)
    _header(ws.cell(top + 1, left), "From \\ To")
    for j, lab in enumerate(df.columns):
        _header(ws.cell(top + 1, left + 1 + j), lab)
    _header(ws.cell(top + 1, left + n + 1), "Total")
    for i, lab in enumerate(df.index):
        r = top + 2 + i
        _header(ws.cell(r, left), lab)
        for j in range(n):
            fill = DIAGONAL_FILL if i == j else DEFAULT_FILL if j == n - 1 else None
            _value(ws.cell(r, left + 1 + j), float(df.iat[i, j]), number_format, fill=fill)
        _value(ws.cell(r, left + n + 1), f"=SUM({L(left + 1)}{r}:{L(left + n)}{r})", "0.00%")


def _matrices_sheet(wb, matrices, model, scenario):
    ws = wb.create_sheet(f"{model} matrices")
    _title(ws, f"{model} - {scenario} scenario - migration matrices",
           "Rows = class at the start, columns = class at the end, D = default. "
           "Left: annual matrix. Right: quarterly matrix applied in the projection (4th root of the annual matrix).")
    n = matrices.n_buckets[model]
    right, top = n + 4, 4
    last = matrices.n_years[model, scenario]
    for y, label in matrices.year_labels(model, scenario).items():
        ws.cell(top, 1, f"Year {label[1:]}").font = _font(bold=True, size=11)
        _write_matrix(ws, top + 1, 1, "Annual", matrices.annual_df(model, scenario, y), "0.00%")
        _write_matrix(ws, top + 1, right, f"Quarterly ({_quarters(y, last)})", matrices.quarterly_df(model, scenario, y), "0.000%")
        top += n + 5
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions[L(n + 3)].width = 3
    ws.column_dimensions[L(right)].width = 10
    for c in range(2, n + 3):
        ws.column_dimensions[L(c)].width = 8.5
    for c in range(right + 1, right + n + 2):   # 3 decimals: 100.000% needs a wider column
        ws.column_dimensions[L(c)].width = 9.5


def _portfolio_sheet(wb, matrices, proj, amortization, reference_date):
    model, scenario = proj.model, proj.scenario
    ws = wb.create_sheet(f"{model} portfolio")
    _title(ws, f"{model} - {scenario} scenario - portfolio vector per quarter",
           f"Amount per class at the end of each quarter. Each quarter the vector migrates with the quarterly matrix, "
           f"then {REPAYMENT[amortization]}. Defaulted amounts stay in D.")
    labels = list(proj.vectors.columns)
    n = len(labels)
    hdr, c_cls = 4, 4
    c_rep, c_total = c_cls + n, c_cls + n + 1
    for j, h in enumerate(["Quarter", "Date", "Matrix used", *labels, "Cumulative repaid", "Total"], start=1):
        _header(ws.cell(hdr, j), h)
    year_labels = matrices.year_labels(model, scenario)
    last = matrices.n_years[model, scenario]
    for t in proj.vectors.index:
        r = hdr + 1 + t
        _value(ws.cell(r, 1), int(t), "0")
        _value(ws.cell(r, 2), (pd.Timestamp(reference_date) + pd.DateOffset(months=3 * int(t))).to_pydatetime(), "yyyy-mm-dd")
        _value(ws.cell(r, 3), "start" if t == 0 else f"Year {year_labels[min((t - 1) // QUARTERS_PER_YEAR + 1, last)][1:]}", "@")
        for k in range(n):
            _value(ws.cell(r, c_cls + k), float(proj.vectors.iat[t, k]), AMOUNT, fill=DEFAULT_FILL if k == n - 1 else None)
        _value(ws.cell(r, c_rep), 0.0 if t == 0 else float(proj.flows.at[t, "cumulative_repaid"]), AMOUNT)
        _value(ws.cell(r, c_total), f"=SUM({L(c_cls)}{r}:{L(c_rep)}{r})", AMOUNT, bold=True)
    ws.cell(hdr + len(proj.vectors) + 2, 1,
            "Total = all classes including D + cumulative repaid; it stays equal to the starting amount every quarter."
            ).font = _font(italic=True, color="595959")
    for c, w in {1: 8, 2: 11, 3: 11, **{c: 12 for c in range(c_cls, c_total + 1)}}.items():
        ws.column_dimensions[L(c)].width = w
    ws.row_dimensions[hdr].height = 30
    ws.freeze_panes = ws.cell(hdr + 1, c_cls)


def export_workbook(path, matrices, portfolio, scenario=DEFAULT_SCENARIO, amortization="bullet", reference_date="2026-10-23"):
    """Write the report for one scenario: per model, its matrices and its portfolio vector per quarter."""
    projections = project_portfolio(portfolio, matrices, scenario, amortization, reference_date)
    wb = Workbook()
    wb.remove(wb.active)
    for model, proj in projections.items():
        _matrices_sheet(wb, matrices, model, scenario)
        _portfolio_sheet(wb, matrices, proj, amortization, reference_date)
    wb.save(path)
    return path
